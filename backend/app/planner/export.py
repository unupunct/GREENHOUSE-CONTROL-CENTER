"""Shopping list exports: CSV, Excel, PDF."""

from __future__ import annotations

import csv
import io

HEADERS = ["Subsystem", "Group", "Item", "Specification", "Qty", "Qty basis / missing inputs",
           "Unit price RON (incl. VAT)", "Price status", "Total RON", "Supplier", "URL",
           "Availability", "Purchased"]


def _rows(data: dict) -> list[list]:
    names = {}
    from app.planner.suppliers import SUPPLIERS
    names = {s["id"]: s["name"] for s in SUPPLIERS}
    rows = []
    for ln in data["lines"]:
        best = next((o for o in ln["offers"] if o["id"] == ln["best_offer_id"]), None)
        basis = ln["qty_basis"] + (f" [needs: {', '.join(ln['qty_missing'])}]" if ln["qty_missing"] else "")
        rows.append([ln["subsystem"], ln["group"], ln["name"], ln["spec"],
                     "" if ln["qty"] is None else ln["qty"], basis,
                     "" if ln["unit_price_ron"] is None else ln["unit_price_ron"],
                     {"verified": "verified", "user_estimate": "your estimate",
                      "none": "no price - requires verification"}[ln["price_status"]],
                     "" if ln["total_ron"] is None else ln["total_ron"],
                     names.get(best["supplier_id"], best["supplier_id"]) if best else "",
                     best["url"] if best else "", best["availability"] if best else "",
                     "yes" if ln["purchased"] else ""])
    return rows


def to_csv(data: dict) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(HEADERS)
    w.writerows(_rows(data))
    b = data["budget"]
    w.writerow([])
    w.writerow(["Total of priced lines (RON)", b["grand_total_ron"], "unpriced lines", b["unpriced_lines"]])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel reads UTF-8


def to_xlsx(data: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Shopping list"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in _rows(data):
        ws.append(r)
    for col, width in zip("ABCDEFGHIJKLM", [22, 12, 44, 44, 6, 36, 14, 22, 12, 24, 40, 28, 10], strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    b = data["budget"]
    sh = wb.create_sheet("Budget by subsystem")
    sh.append(["Subsystem", "Lines", "Priced lines", "Needs input", "Priced total RON"])
    for name, v in b["by_subsystem"].items():
        sh.append([name, v["lines"], v["priced_lines"], v["needs_input_lines"], v["priced_total_ron"]])
    sh.append([])
    sh.append(["TOTAL (priced lines only)", b["lines"], b["priced_lines"], b["needs_input_lines"], b["grand_total_ron"]])
    sh.append(["of which verified", "", "", "", b["of_which_verified_ron"]])
    sh.append(["of which your estimates", "", "", "", b["of_which_user_estimate_ron"]])
    sh.append([])
    sh.append(["Tier comparison (priced lines only)"])
    for t, v in data["tier_totals_ron"].items():
        sh.append([t, "", "", "", v])
    sc = wb.create_sheet("Supplier comparison")
    sc.append(["Supplier", "Items covered", "Coverage %", "Total RON"])
    for s in data["supplier_comparison"]:
        sc.append([s["supplier"], s["items_covered"], s["coverage_pct"], s["total_ron"]])
    note = wb.create_sheet("Notes")
    note.append([data["label"]])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _p(text) -> str:
    """Core PDF fonts are Latin-1 only; replace what they cannot show."""
    s = str(text).replace("≥", ">=").replace("≤", "<=").replace("→", "->").replace("÷", "/")
    s = s.replace("₂", "2").replace("–", "-").replace("’", "'")
    return s.encode("latin-1", "replace").decode("latin-1")


def to_pdf(data: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=7, leading=8.5)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
    story = [Paragraph(_p(f"Greenhouse shopping list - tier: {data['tier']}"), styles["Title"]),
             Paragraph(_p(data["label"]), styles["BodyText"]), Spacer(1, 8)]
    head = ["Item", "Qty", "Unit RON", "Total RON", "Price status", "Supplier / URL"]
    from app.planner.suppliers import SUPPLIERS
    names = {s["id"]: s["name"] for s in SUPPLIERS}
    rows = [head]
    for ln in data["lines"]:
        best = next((o for o in ln["offers"] if o["id"] == ln["best_offer_id"]), None)
        qty = "?" if ln["qty"] is None else str(ln["qty"])
        status = {"verified": "verified", "user_estimate": "your estimate", "none": "needs verification"}[ln["price_status"]]
        sup = _p(f"{names.get(best['supplier_id'], best['supplier_id'])} {best['url']}") if best else ""
        rows.append([Paragraph(_p(f"<b>{ln['name']}</b> - {ln['spec']}"), small), qty,
                     "" if ln["unit_price_ron"] is None else f"{ln['unit_price_ron']:.2f}",
                     "" if ln["total_ron"] is None else f"{ln['total_ron']:.2f}", status, Paragraph(sup, small)])
    t = Table(rows, repeatRows=1, colWidths=[300, 30, 55, 60, 80, 270])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2e8b57")),
                           ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTSIZE", (0, 0), (-1, -1), 7),
                           ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(t)
    b = data["budget"]
    story += [Spacer(1, 10), Paragraph(_p(
        f"<b>Total of priced lines: {b['grand_total_ron']:.2f} RON</b> ({b['priced_lines']} priced, "
        f"{b['unpriced_lines']} without price, {b['needs_input_lines']} quantities need input)."), styles["BodyText"])]
    doc.build(story)
    return buf.getvalue()
