import csv
import io

import pytest
from openpyxl import load_workbook

from app.planner import sizing


def _opts(client, **kw):
    o = client.get("/api/planner/options").json()["options"]
    o.update(kw)
    assert client.put("/api/planner/options", json=o).status_code == 200


def _cfg(client, **kw):
    c = client.get("/api/config").json()
    c.update(kw)
    assert client.put("/api/config", json=c).status_code == 200


def test_cable_known_values():
    r = sizing.cable(20, 5, 24, 3)
    assert r["suggested_mm2"] == 6 and abs(r["min_cross_section_mm2"] - 4.861) < 0.001
    assert abs(r["voltage_drop_v"] - 0.583) < 0.001 and abs(r["voltage_drop_pct"] - 2.43) < 0.01
    assert "ampacity" in r["note"]


def test_plan_without_inputs_never_invents_numbers(client):
    p = client.get("/api/planner/plan").json()
    assert p["greenhouse"]["area_m2"] == 200
    assert p["climate_stations"]["count"] == 2
    assert p["areas"]["soil_m2"] == 100
    assert p["irrigation"]["soil_zones"] == 2
    v = p["ventilation"]
    assert v["status"] == "needs_input" and v["fans"] is None and v["volume_m3"] is None
    assert any("height" in m for m in v["missing"])
    assert p["lighting"]["status"] == "needs_input" and "fixtures" not in p["lighting"]
    assert "Raspberry Pi 5" in p["controller"]["recommended"]


def test_ventilation_needs_height_then_calculates(client):
    _opts(client, ach_target=30, fan_flow_m3_h=10000)
    assert client.get("/api/planner/plan").json()["ventilation"]["fans"] is None  # still no height
    _cfg(client, height_m=3.0)
    v = client.get("/api/planner/plan").json()["ventilation"]
    assert v["volume_m3"] == 600 and v["required_airflow_m3_h"] == 18000
    assert v["fans"] == 3 and v["status"] == "final"
    assert "underestimated" in v["note"]
    _cfg(client, ridge_height_m=5.0)
    assert client.get("/api/planner/plan").json()["ventilation"]["volume_m3"] == 800  # 200 * (3+1)


def test_lighting_dli_ppfd(client):
    _opts(client, dli_target=20, natural_dli=10, photoperiod_h=16, fixture_ppf=1000)
    li = client.get("/api/planner/plan").json()["lighting"]
    assert abs(li["supplemental_ppfd_umol_m2_s"] - 173.6) < 0.1
    assert li["fixtures"] == 44 and "lux" in li["note"]
    _opts(client, dli_target=20, natural_dli=25, photoperiod_h=16, fixture_ppf=1000)
    assert client.get("/api/planner/plan").json()["lighting"]["fixtures"] == 0


def test_irrigation_calculator(client):
    _opts(client, drip_line_total_m=400, dripper_spacing_m=0.5, dripper_flow_l_h=2,
          pump_flow_l_min=50, water_l_per_m2_day=3, autonomy_days=2)
    r = client.get("/api/planner/plan").json()["irrigation"]
    assert r["drippers"] == 800 and r["all_drippers_flow_l_h"] == 1600
    assert r["flow_per_soil_zone_l_h"] == 800
    assert r["daily_water_l_from_demand"] == 600 and r["reservoir_min_l"] == 1200
    assert r["zones_simultaneously_supported"] == 3  # 50 / (800/60)
    assert r["status"] == "final"


def test_electrical_loads(client):
    _opts(client, loads=[{"name": "valve", "power_w": 10, "voltage": "24VDC", "qty": 4},
                         {"name": "fan", "power_w": 400, "voltage": "230VAC", "qty": 2}])
    e = client.get("/api/planner/plan").json()["electrical"]
    assert e["psu_24v_min_w"] == 50
    assert e["dc_fuse_suggestions"][0]["suggested_fuse_a"] == 1  # 10 W / 24 V = 0.42 A, x1.25 = 0.52 A
    assert "electrician" in e["mains_note"]
    assert all("rcd" not in k for k in e)  # no mains protective device sizing offered


def test_preview_does_not_save(client):
    o = client.get("/api/planner/options").json()["options"]
    o["ach_target"] = 60
    assert client.post("/api/planner/plan", json=o).status_code == 200
    assert client.get("/api/planner/options").json()["options"]["ach_target"] is None


def test_shopping_prices_only_from_offers(client):
    s = client.get("/api/shopping").json()
    priced = [ln for ln in s["lines"] if ln["unit_price_ron"] is not None]
    assert [ln["id"] for ln in priced] == ["pi5_4gb"]  # the one verified, sourced offer
    pi = priced[0]
    assert pi["unit_price_ron"] == 639 and pi["price_status"] == "verified"
    assert pi["offers"][0]["url"].startswith("https://www.optimusdigital.ro/")
    assert s["budget"]["grand_total_ron"] == 639
    assert s["budget"]["unpriced_lines"] == s["budget"]["lines"] - 1
    assert all(ln["price_status"] == "none" for ln in s["lines"] if ln["id"] != "pi5_4gb")
    assert "not estimate" in s["label"]


def test_tiers_and_groups(client):
    low = client.get("/api/shopping", params={"tier": "low"}).json()["lines"]
    pro = client.get("/api/shopping", params={"tier": "professional"}).json()["lines"]
    assert len(low) < len(pro)
    ids = {ln["id"] for ln in low}
    assert "estop" in ids and "float_sw" in ids  # hardware protection is in every tier
    no_opt = client.get("/api/shopping", params={"tier": "professional", "include_optional": False}).json()
    assert all(ln["group"] != "optional" for ln in no_opt["lines"])
    groups = {ln["group"] for ln in pro}
    assert {"essential", "spare", "consumable", "tool", "installation"} <= groups


def test_single_method_greenhouse_changes_bom(client):
    _cfg(client, growing_methods=["soil"])
    ids = {ln["id"] for ln in client.get("/api/shopping").json()["lines"]}
    assert "ph_probe" not in ids and "soil_probe" in ids
    _cfg(client, growing_methods=["hydroponic"])
    ids = {ln["id"] for ln in client.get("/api/shopping").json()["lines"]}
    assert "ph_probe" in ids and "soil_probe" not in ids


def test_offers_vat_and_purchased(client):
    bad = {"item_id": "nope", "supplier_id": "x", "price_ron": 10}
    assert client.post("/api/shopping/offers", json=bad).status_code == 422
    ok = {"item_id": "fuses", "supplier_id": "soselectronic", "price_ron": 100, "includes_vat": False,
          "url": "https://example.test/p", "status": "user_estimate"}
    oid = client.post("/api/shopping/offers", json=ok).json()["id"]
    s = client.get("/api/shopping").json()
    line = next(ln for ln in s["lines"] if ln["id"] == "fuses")
    assert line["unit_price_ron"] == 121 and line["price_status"] == "user_estimate"
    assert s["budget"]["grand_total_ron"] == 760
    assert s["budget"]["of_which_verified_ron"] == 639 and s["budget"]["of_which_user_estimate_ron"] == 121
    assert {c["supplier_id"] for c in s["supplier_comparison"]} == {"optimusdigital", "soselectronic"}
    assert client.post("/api/shopping/offers", json={**ok, "url": "javascript:x"}).status_code == 422
    # a verified offer beats a cheaper estimate
    cheap = {**ok, "price_ron": 1, "supplier_id": "mine"}
    client.post("/api/shopping/offers", json=cheap)
    client.post("/api/shopping/offers", json={**ok, "price_ron": 500, "status": "verified"})
    line = next(ln for ln in client.get("/api/shopping").json()["lines"] if ln["id"] == "fuses")
    assert line["price_status"] == "verified"
    assert client.delete(f"/api/shopping/offers/{oid}").status_code == 200
    assert client.delete("/api/shopping/offers/99999").status_code == 404
    client.put("/api/shopping/purchased/estop", json={"purchased": True})
    line = next(ln for ln in client.get("/api/shopping").json()["lines"] if ln["id"] == "estop")
    assert line["purchased"] is True


def test_exports_are_valid_files(client):
    c = client.get("/api/shopping/export.csv")
    rows = list(csv.reader(io.StringIO(c.content.decode("utf-8-sig"))))
    assert rows[0][0] == "Subsystem" and any("Raspberry Pi 5" in r[2] for r in rows if len(r) > 2)
    x = client.get("/api/shopping/export.xlsx")
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Shopping list", "Budget by subsystem", "Supplier comparison", "Notes"]
    p = client.get("/api/shopping/export.pdf")
    assert p.content[:5] == b"%PDF-" and len(p.content) > 2000
    assert client.get("/api/shopping/export.doc").status_code == 422


def test_suppliers_do_not_claim_unverified_urls(client):
    d = client.get("/api/shopping/suppliers").json()
    assert "Dedeman" in d["unverified_targets"]
    assert all(s["url"].startswith("https://") and s["verified"] for s in d["suppliers"])


def test_wiring_diagrams_and_tables(client):
    d = client.get("/api/wiring/diagrams").json()
    assert d["architecture"].startswith("flowchart") and "USB-RS485" in d["architecture"]
    assert "Hardware protection" in d["power"] and "Float switch" in d["power"]
    assert "No RS485 bus" in d["rs485"]
    client.put("/api/hardware/buses", json=[{"id": "rs485_1", "port": "COM9"}])
    for i, reg in enumerate((0, 1)):
        client.post("/api/sensors", json={"id": f"mb_{i}", "name": "n", "kind": "temperature", "zone": "soil",
                                          "unit": "C", "driver": "modbus_rtu",
                                          "address": {"bus": "rs485_1", "slave": 4, "register": reg}})
    d = client.get("/api/wiring/diagrams").json()
    assert "Slave 4" in d["rs485"] and "mb_0, mb_1" in d["rs485"] and "termination" in d["rs485"]
    t = client.get("/api/wiring/modbus").json()
    assert [r["register"] for r in t] == [0, 1] and t[0]["sensors_on_this_address"] == 2


def test_gpio_assignments_validated(client):
    g = client.get("/api/wiring/gpio").json()
    assert {p["bcm"] for p in g["fixed"]} == {2, 3, 4}
    ok = [{"bcm": 17, "function": "Valve 1 driver input", "device": "driver ch1"}]
    assert client.put("/api/wiring/gpio", json=ok).status_code == 200
    assert client.get("/api/wiring/gpio").json()["assigned"][0]["bcm"] == 17
    for bad in ([{"bcm": 2, "function": "x"}], [{"bcm": 99, "function": "x"}],
                [{"bcm": 5, "function": "a"}, {"bcm": 5, "function": "b"}]):
        assert client.put("/api/wiring/gpio", json=bad).status_code == 422


@pytest.mark.parametrize("path", ["/api/planner/plan", "/api/wiring/diagrams"])
def test_endpoints_ok_on_fresh_install(client, path):
    assert client.get(path).status_code == 200
